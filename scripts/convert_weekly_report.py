from __future__ import annotations

import html
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


PROJECT_DIR = Path(__file__).resolve().parents[1]
REPORTS_DIR = PROJECT_DIR / "reports" / "Pharma"
REPORT_JSON = REPORTS_DIR / "weekly_report_20260703.json"
REPORT_MD = REPORTS_DIR / "weekly_report_20260703.md"
OUT_HTML = REPORTS_DIR / "weekly_report_20260703.html"
OUT_PDF = REPORTS_DIR / "weekly_report_20260703.pdf"
OUT_DOCX = REPORTS_DIR / "weekly_report_20260703.docx"


def load_markdown() -> str:
    if REPORT_JSON.exists():
        data = json.loads(REPORT_JSON.read_text(encoding="utf-8"))
        text = data.get("markdown_report", "")
    else:
        text = REPORT_MD.read_text(encoding="utf-8")
    start = text.find("# Weekly Pharma News Report")
    if start > 0:
        text = text[start:]
    return text.strip()


def clean_inline_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def parse_markdown_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [clean_inline_markdown(cell.strip()) for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        i += 1
    return rows, i


def markdown_to_html(md: str) -> str:
    lines = md.splitlines()
    body: list[str] = []
    in_ul = False
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            if in_ul:
                body.append("</ul>")
                in_ul = False
            i += 1
            continue
        if stripped.startswith("|"):
            if in_ul:
                body.append("</ul>")
                in_ul = False
            rows, i = parse_markdown_table(lines, i)
            if rows:
                body.append("<table>")
                for row_index, row in enumerate(rows):
                    tag = "th" if row_index == 0 else "td"
                    body.append("<tr>" + "".join(f"<{tag}>{html.escape(cell)}</{tag}>" for cell in row) + "</tr>")
                body.append("</table>")
            continue
        if stripped.startswith("#"):
            if in_ul:
                body.append("</ul>")
                in_ul = False
            level = min(len(stripped) - len(stripped.lstrip("#")), 3)
            text = clean_inline_markdown(stripped.lstrip("#").strip())
            body.append(f"<h{level}>{html.escape(text)}</h{level}>")
        elif stripped.startswith("- "):
            if not in_ul:
                body.append("<ul>")
                in_ul = True
            body.append(f"<li>{html.escape(clean_inline_markdown(stripped[2:]))}</li>")
        elif stripped == "---":
            if in_ul:
                body.append("</ul>")
                in_ul = False
            body.append("<hr>")
        else:
            body.append(f"<p>{html.escape(clean_inline_markdown(stripped))}</p>")
        i += 1
    if in_ul:
        body.append("</ul>")

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <title>Weekly Pharma News Report</title>
  <style>
    body {{
      margin: 0;
      background: #f4f6f8;
      color: #18202a;
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "PingFang SC", "Hiragino Sans GB", Arial, sans-serif;
      line-height: 1.62;
    }}
    main {{
      max-width: 920px;
      margin: 32px auto;
      background: #fff;
      padding: 48px 56px;
      box-shadow: 0 10px 35px rgba(20, 34, 54, .12);
      border-radius: 8px;
    }}
    h1 {{ font-size: 30px; margin: 0 0 20px; color: #0f2d4a; }}
    h2 {{ font-size: 21px; margin: 34px 0 12px; color: #174f7a; border-bottom: 1px solid #d9e2ec; padding-bottom: 6px; }}
    h3 {{ font-size: 17px; margin: 24px 0 8px; color: #263746; }}
    p {{ margin: 8px 0; }}
    li {{ margin: 6px 0; }}
    table {{ border-collapse: collapse; width: 100%; margin: 14px 0 24px; font-size: 13px; }}
    th, td {{ border: 1px solid #d8dee6; padding: 8px 10px; vertical-align: top; }}
    th {{ background: #edf4fa; text-align: left; }}
    hr {{ border: 0; border-top: 1px solid #d8dee6; margin: 28px 0; }}
    @media print {{
      body {{ background: #fff; }}
      main {{ margin: 0; box-shadow: none; border-radius: 0; max-width: none; }}
    }}
  </style>
</head>
<body>
<main>
{chr(10).join(body)}
</main>
</body>
</html>
"""


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_docx_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(clean_inline_markdown(text))
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")


def markdown_to_docx(md: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_markdown_table(lines, i)
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx, cell_text in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_text
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(8.5)
                                run.font.name = "Arial"
                                run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
                        if r_idx == 0:
                            set_cell_shading(cell, "EAF2F8")
                doc.add_paragraph()
            continue
        if stripped.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(clean_inline_markdown(stripped[2:]))
            run.bold = True
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(15, 45, 74)
            run.font.name = "Arial"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        elif stripped.startswith("## "):
            add_docx_paragraph(doc, stripped[3:], "Heading 1")
        elif stripped.startswith("### "):
            add_docx_paragraph(doc, stripped[4:], "Heading 2")
        elif stripped.startswith("- "):
            add_docx_paragraph(doc, stripped[2:], "List Bullet")
        elif stripped == "---":
            doc.add_paragraph()
        else:
            add_docx_paragraph(doc, stripped)
        i += 1
    doc.save(OUT_DOCX)


def pdf_styles():
    pdfmetrics.registerFont(TTFont("STHeiti", "/System/Library/Fonts/STHeiti Light.ttc"))
    styles = getSampleStyleSheet()
    base = {
        "fontName": "STHeiti",
        "alignment": TA_LEFT,
    }
    return {
        "h1": ParagraphStyle("H1", parent=styles["Heading1"], fontSize=20, leading=25, textColor=colors.HexColor("#0f2d4a"), spaceAfter=12, **base),
        "h2": ParagraphStyle("H2", parent=styles["Heading2"], fontSize=15, leading=19, textColor=colors.HexColor("#174f7a"), spaceBefore=14, spaceAfter=8, **base),
        "h3": ParagraphStyle("H3", parent=styles["Heading3"], fontSize=12.5, leading=16, textColor=colors.HexColor("#263746"), spaceBefore=10, spaceAfter=5, **base),
        "body": ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9.5, leading=14.5, spaceAfter=6, **base),
        "bullet": ParagraphStyle("Bullet", parent=styles["BodyText"], fontSize=9.5, leading=14, leftIndent=12, **base),
        "table": ParagraphStyle("Table", parent=styles["BodyText"], fontSize=7.2, leading=9.5, **base),
    }


def markdown_to_pdf(md: str) -> None:
    style = pdf_styles()
    story = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_markdown_table(lines, i)
            if rows:
                data = [[Paragraph(html.escape(cell), style["table"]) for cell in row] for row in rows]
                col_count = max(len(row) for row in rows)
                table = Table(data, colWidths=[7.0 * inch / col_count] * col_count, repeatRows=1)
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF2F8")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C8D2DC")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]))
                story.append(table)
                story.append(Spacer(1, 10))
            continue
        if stripped.startswith("# "):
            story.append(Paragraph(html.escape(clean_inline_markdown(stripped[2:])), style["h1"]))
        elif stripped.startswith("## "):
            story.append(Paragraph(html.escape(clean_inline_markdown(stripped[3:])), style["h2"]))
        elif stripped.startswith("### "):
            story.append(Paragraph(html.escape(clean_inline_markdown(stripped[4:])), style["h3"]))
        elif stripped.startswith("- "):
            story.append(ListFlowable([ListItem(Paragraph(html.escape(clean_inline_markdown(stripped[2:])), style["bullet"]))], bulletType="bullet"))
        elif stripped == "---":
            story.append(Spacer(1, 8))
        else:
            story.append(Paragraph(html.escape(clean_inline_markdown(stripped)), style["body"]))
        i += 1

    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=LETTER,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title="Weekly Pharma News Report",
    )
    doc.build(story)


def main() -> None:
    md = load_markdown()
    OUT_HTML.write_text(markdown_to_html(md), encoding="utf-8")
    markdown_to_docx(md)
    markdown_to_pdf(md)
    print(f"HTML: {OUT_HTML}")
    print(f"DOCX: {OUT_DOCX}")
    print(f"PDF: {OUT_PDF}")


if __name__ == "__main__":
    main()
