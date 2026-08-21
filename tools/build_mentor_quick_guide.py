from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "mentor_codex_wewerss_quick_guide.docx"

BLUE = "245B78"
LIGHT_BLUE = "EAF3F8"
LIGHT_GRAY = "F4F6F8"
GREEN = "2E6B57"
GOLD = "8A6500"
TEXT = "24313A"
MUTED = "667680"


def set_font(run, name="Arial Unicode MS", size=11, bold=False, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    ppr.append(borders)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_font(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_font(p.add_run(text))
    return p


def add_prompt(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.15
    shade_paragraph(p, LIGHT_GRAY)
    for i, line in enumerate(lines):
        if i:
            p.add_run("\n")
        set_font(p.add_run(line), name="Arial Unicode MS", size=10, color=TEXT)
    return p


def add_success(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    shade_paragraph(p, "EDF6F1")
    set_font(p.add_run("✓ 成功标志："), bold=True, color=GREEN)
    set_font(p.add_run(text), color=TEXT)


def add_warning(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    shade_paragraph(p, "FFF8E6")
    set_font(p.add_run("重要提醒："), bold=True, color=GOLD)
    set_font(p.add_run(text), color=TEXT)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after in [
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 14, 7),
    ]:
        st = doc.styles[name]
        st.font.name = "Arial Unicode MS"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(BLUE)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        st = doc.styles[name]
        st.font.name = "Arial Unicode MS"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, size=9, color=MUTED)


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("MENTOR 操作指南  |  WeWe RSS × Codex × DeepSeek"), size=8.5, color=MUTED)
    set_page_number(section.footer.paragraphs[0])

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(5)
    set_font(kicker.add_run("极简操作版"), size=10, bold=True, color=GREEN)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    set_font(title.add_run("Codex 公众号月报使用说明"), size=25, bold=True, color=BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    set_font(subtitle.add_run("适用于医药与医疗器械月报｜按步骤复制指令即可"), size=11.5, color=MUTED)

    add_warning(doc, "每完成一步就停止，确认结果后再进入下一步。不要提前爬取文章或调用 DeepSeek。")

    add_heading(doc, "1｜连接 WeWe RSS")
    add_body(doc, "复制给 Codex：", bold_lead="复制给 Codex：")
    add_prompt(doc, [
        "请启动并检查本地 WeWe RSS，然后打开：",
        "http://localhost:8001",
        "确认页面可以正常访问后停止，不要爬取文章，也不要调用 DeepSeek。",
    ])
    add_success(doc, "浏览器打开后进入 WeWe RSS 页面。")

    add_heading(doc, "2｜登录微信并添加公众号")
    add_body(doc, "这一步需要 mentor 手动操作：", bold_lead="这一步需要 mentor 手动操作：")
    for item in [
        "进入“账号管理”，删除已经失效的账号。",
        "用微信扫码登录。",
        "回到“公众号”页面。",
        "用微信打开想添加的公众号文章。",
        "选择“在浏览器打开”（推荐 Google Chrome）。",
        "复制文章链接，粘贴到 WeWe RSS。",
        "等待 WeWe RSS 自动识别并添加该公众号。",
    ]:
        add_number(doc, item)
    add_success(doc, "新公众号出现在 WeWe RSS 的公众号列表中。")

    add_heading(doc, "3｜爬取指定日期的文章")
    add_body(doc, "复制给 Codex，并替换开始和结束日期：", bold_lead="复制给 Codex，并替换开始和结束日期：")
    add_prompt(doc, [
        "请从 WeWe RSS 爬取以下日期范围内的公众号文章：",
        "开始日期：YYYY-MM-DD",
        "结束日期：YYYY-MM-DD",
        "",
        "请根据公众号来源，将文章标注为“医药”或“医疗器械”。",
        "爬取完成后检查日期、正文、乱码和无效网页内容，并分别统计两类文章数量。",
        "检查完成后停止，不要调用 DeepSeek API。",
    ])
    add_body(doc, "检查项目：", bold_lead="检查项目：")
    for item in [
        "文章日期在指定范围内。",
        "正文不为空，并且是正常、可阅读的中文字符。",
        "没有乱码、网页代码或无效内容。",
        "已分别统计医药和医疗器械文章数量。",
    ]:
        add_bullet(doc, item)
    add_success(doc, "每篇文章都有日期、有效正文和正确的来源分类。")

    add_heading(doc, "4｜接入 DeepSeek 开始分析")
    add_body(doc, "复制给 Codex：", bold_lead="复制给 Codex：")
    add_prompt(doc, [
        "请先告诉我医药和医疗器械分别有多少篇有效文章，以及预计需要运行哪些分析。",
        "等我确认后，再依次运行以下四项：",
        "1. 医药 Article Analysis",
        "2. 医疗器械 Article Analysis",
        "3. 医药 Monthly Analysis",
        "4. 医疗器械 Monthly Analysis",
        "",
        "两条业务线必须使用独立的输入、分析结果和月报文件，不能混在一起。",
        "每完成一项后汇报结果并停止，等我确认后再进行下一项。",
    ])

    add_heading(doc, "四次分析分别做什么？", level=2)
    for label, detail in [
        ("医药 Article Analysis", "对所有医药文章逐篇打分和总结。"),
        ("医疗器械 Article Analysis", "对所有医疗器械文章逐篇打分和总结。"),
        ("医药 Monthly Analysis", "根据分数筛选 Target Articles，并撰写医药月报。"),
        ("医疗器械 Monthly Analysis", "根据分数筛选 Target Articles，并撰写医疗器械月报。"),
    ]:
        p = doc.add_paragraph()
        set_font(p.add_run(label + "："), bold=True, color=BLUE)
        set_font(p.add_run(detail))

    add_success(doc, "最终得到 1 份医药月报和 1 份医疗器械月报；合计运行 4 次 DeepSeek Analysis。")
    add_warning(doc, "调用 DeepSeek 前先检查文章数量和正文质量。不要随意使用 --force，以免重复分析并产生额外 API 费用。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
