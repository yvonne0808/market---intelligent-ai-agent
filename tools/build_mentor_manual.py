from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/mentor_market_intelligence_manual.md"
OUT_DIR = ROOT / "artifacts"
OUTPUT = OUT_DIR / "mentor_market_intelligence_manual.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6770"
CODE_FILL = "E8EEF5"
CAUTION_FILL = "FDECEC"
SUCCESS_FILL = "EAF6EE"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths=(2700, 6660)) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = "Hiragino Sans GB" if name == "Calibri" else name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    rfonts.set(qn("w:cs"), "Hiragino Sans GB")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para(paragraph, before=0, after=6, line=1.25, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def add_code(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, (9360,))
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    p = cell.paragraphs[0]
    set_para(p, before=0, after=0, line=1.05)
    for idx, line in enumerate(text.rstrip("\n").splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, name="Menlo", size=8.5, color=INK)
    spacer = doc.add_paragraph()
    set_para(spacer, after=4)


def add_callout(doc: Document, label: str, text: str, fill: str, color: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, (9360,))
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_para(p, before=0, after=0, line=1.15)
    run = p.add_run(label + "：")
    set_font(run, size=10.5, color=color, bold=True)
    run = p.add_run(text)
    set_font(run, size=10.5, color=INK)
    spacer = doc.add_paragraph()
    set_para(spacer, after=4)


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = len(rows[0])
    widths = (2700, 6660) if cols == 2 else tuple([9360 // cols] * cols)
    table = doc.add_table(rows=len(rows), cols=cols)
    set_table_geometry(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            set_cell_shading(cell, "E8EEF5" if r_idx == 0 else WHITE)
            p = cell.paragraphs[0]
            set_para(p, before=0, after=0, line=1.15)
            run = p.add_run(re.sub(r"\*\*(.*?)\*\*", r"\1", value.strip()))
            set_font(run, size=9.2 if cols > 2 else 10, color=INK, bold=(r_idx == 0))
    spacer = doc.add_paragraph()
    set_para(spacer, after=4)


def add_heading(doc: Document, text: str, level: int):
    if level == 1 and (text.startswith("Step") or "｜" in text):
        doc.add_page_break()
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(re.sub(r"\*\*(.*?)\*\*", r"\1", text))
    return p


def add_inline_runs(paragraph, text: str):
    for part in re.split(r"(`[^`]+`|\*\*.*?\*\*)", text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Menlo", size=9.5, color=INK)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def parse_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def build():
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Hiragino Sans GB"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Hiragino Sans GB")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Hiragino Sans GB")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Hiragino Sans GB")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Hiragino Sans GB"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Hiragino Sans GB")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Hiragino Sans GB")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        style._element.rPr.rFonts.set(qn("w:cs"), "Hiragino Sans GB")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    set_para(header, after=0)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("Market Intelligence · Mentor Guide")
    set_font(run, size=9, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    set_para(footer, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    run = footer.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    add_page_field(footer)

    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []
    paragraph_lines: list[str] = []
    first_title_done = False

    def flush_paragraph():
        nonlocal paragraph_lines
        if not paragraph_lines:
            return
        text = " ".join(line.strip() for line in paragraph_lines).strip()
        if text:
            p = doc.add_paragraph()
            set_para(p)
            add_inline_runs(p, text)
        paragraph_lines = []

    def flush_table():
        nonlocal table_lines
        if table_lines:
            rows = [parse_table_row(line) for line in table_lines if not re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$", line)]
            add_table(doc, rows)
        table_lines = []

    for line in source_lines:
        if line.startswith("```"):
            flush_paragraph()
            flush_table()
            if in_code:
                add_code(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("|"):
            flush_paragraph()
            table_lines.append(line)
            continue
        if table_lines and not line.startswith("|"):
            flush_table()
        if not line.strip():
            flush_paragraph()
            continue
        if line.strip() == "---":
            flush_paragraph()
            p = doc.add_paragraph()
            set_para(p, before=2, after=8)
            pPr = p._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), BLUE)
            border.append(bottom)
            pPr.append(border)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            flush_paragraph()
            add_heading(doc, heading.group(2).strip(), len(heading.group(1)))
            first_title_done = True
            continue
        if line.startswith(">"):
            flush_paragraph()
            text = line[1:].strip()
            if "停止点" in text:
                add_callout(doc, "停止点", text.replace("**", ""), CAUTION_FILL, "9B1C1C")
            else:
                add_callout(doc, "提示", text.replace("**", ""), CODE_FILL, DARK_BLUE)
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet or numbered:
            flush_paragraph()
            style_name = "List Bullet" if bullet else "List Number"
            p = doc.add_paragraph(style=style_name)
            set_para(p, after=4, line=1.25)
            text = (bullet or numbered).group(1)
            add_inline_runs(p, text)
            continue
        paragraph_lines.append(line)
    flush_paragraph()
    flush_table()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
